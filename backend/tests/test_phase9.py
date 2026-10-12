"""فاز ۹ - بند ۱۷ سند: آزمون دسترسی، اکسل، هم‌زمانی، فعال‌سازی، رخداد"""
import io

from conftest import PASSWORD, bearer, login

API = "/api/v1"


# ---------- دسترسی ----------

def test_wrong_password_401(client):
    r = client.post(f"{API}/auth/login", json={"username": "t_admin", "password": "WrongPass!999"})
    assert r.status_code == 401


def test_unauthenticated_401(client):
    assert client.get(f"{API}/cases").status_code == 401


def test_system_admin_cannot_list_cases(client, env):
    t = login(client, "t_admin")
    assert client.get(f"{API}/cases", headers=bearer(t)).status_code == 403


def test_counselor_b_cannot_see_case_of_a(client, env):
    tb = login(client, "t_counsb")
    cid = str(env["case"].id)
    assert client.get(f"{API}/cases/{cid}", headers=bearer(tb)).status_code == 404


def test_counselor_a_sees_own_case(client, env):
    ta = login(client, "t_counsa")
    cid = str(env["case"].id)
    r = client.get(f"{API}/cases/{cid}", headers=bearer(ta))
    assert r.status_code == 200
    assert r.json()["case_number"] == "RS-TEST-000001"


def test_staff_cannot_access_assessments(client, env):
    cid = str(env["case"].id)
    for u in ("t_manager", "t_reception"):
        t = login(client, u)
        assert client.get(f"{API}/cases/{cid}/assessments", headers=bearer(t)).status_code == 404


def test_applicant_sees_only_own_summary(client, env):
    t = login(client, "t_appa")
    r = client.get(f"{API}/me/cases", headers=bearer(t))
    assert r.status_code == 200
    items = r.json()["items"]
    assert len(items) == 1
    assert items[0]["case_number"] == "RS-TEST-000001"
    assert "diagnosis" not in items[0]


def test_applicant_cannot_access_admin_apis(client):
    t = login(client, "t_appa")
    assert client.get(f"{API}/users", headers=bearer(t)).status_code == 403
    assert client.get(f"{API}/reports/summary", headers=bearer(t)).status_code == 403


def test_draft_recommendation_hidden_from_applicant(client, env):
    ta = login(client, "t_counsa")
    cid = str(env["case"].id)
    r = client.post(f"{API}/cases/{cid}/recommendations", headers=bearer(ta),
                    json={"content": "توصیه آزمایشی پیش‌نویس"})
    assert r.status_code == 201
    rec_id = r.json()["id"]
    tp = login(client, "t_appa")
    items = client.get(f"{API}/me/recommendations", headers=bearer(tp)).json()["items"]
    assert all(i["id"] != rec_id for i in items)
    r2 = client.post(f"{API}/recommendations/{rec_id}/publish", headers=bearer(ta))
    assert r2.status_code == 200
    items2 = client.get(f"{API}/me/recommendations", headers=bearer(tp)).json()["items"]
    assert any(i["id"] == rec_id for i in items2)


def test_disabled_user_cannot_login(client, env):
    r = client.post(f"{API}/auth/login", json={"username": "t_disabled", "password": PASSWORD})
    assert r.status_code == 401


def test_activation_flow(client, env):
    ta = login(client, "t_admin")
    uid = str(env["pending_user"].id)
    r = client.post(f"{API}/users/{uid}/activation-token", headers=bearer(ta))
    assert r.status_code == 200
    token = r.json()["token"]
    r2 = client.post(f"{API}/auth/activate", json={"token": token, "new_password": "NewPass!4567"})
    assert r2.status_code == 200
    r3 = client.post(f"{API}/auth/login", json={"username": "t_pending", "password": "NewPass!4567"})
    assert r3.status_code == 200


# ---------- اکسل (بند ۸) ----------

APPLICANT_HEADERS = ["external_id", "employee_code", "first_name", "last_name", "national_id",
                     "birth_date", "gender", "marital_status", "education", "occupation",
                     "mobile", "phone", "unit_code", "service_location_code", "province_code",
                     "county_code", "city", "address", "referral_source", "referral_description",
                     "create_account", "username"]


def _applicant_file(rows):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(APPLICANT_HEADERS)
    for row in rows:
        ws.append([row.get(h) for h in APPLICANT_HEADERS])
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_template_download(client, env):
    t = login(client, "t_reception")
    r = client.get(f"{API}/imports/templates/applicants", headers=bearer(t))
    assert r.status_code == 200
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[0] == "داده"
    assert wb["داده"]["A1"].value == "external_id"


def test_excel_import_flow_and_idempotency(client, env):
    t = login(client, "t_reception")
    rows = [
        {"external_id": "T-XL-1", "first_name": "الف", "last_name": "یک",
         "national_id": "0098765432", "birth_date": "1370/03/21"},
        {"external_id": "T-XL-2", "first_name": "ب", "last_name": "دو", "employee_code": "9002"},
        {"external_id": "T-XL-3", "first_name": "ج", "last_name": "سه", "national_id": "12345"},
    ]
    r = client.post(f"{API}/imports/applicants/preview", headers=bearer(t),
                    files={"file": ("t.xlsx", _applicant_file(rows), XLSX_MIME)})
    assert r.status_code == 200, r.text
    summary = r.json()["summary"]
    assert summary["valid"] == 2 and summary["error"] == 1
    rc = client.post(f"{API}/imports/{r.json()['job_id']}/commit", headers=bearer(t),
                     json={"duplicate_strategy": "skip"})
    assert rc.status_code == 200, rc.text
    rep = rc.json()
    assert rep["created"] == 2 and rep["errors"] == 1
    # ثبت دوباره همان داده → تکراری، بدون رکورد مضاعف (بند ۸.۵)
    r2 = client.post(f"{API}/imports/applicants/preview", headers=bearer(t),
                     files={"file": ("t2.xlsx", _applicant_file(rows), XLSX_MIME)})
    assert r2.json()["summary"]["duplicate"] == 2
    rc2 = client.post(f"{API}/imports/{r2.json()['job_id']}/commit", headers=bearer(t),
                      json={"duplicate_strategy": "skip"})
    assert rc2.json()["created"] == 0 and rc2.json()["skipped"] == 2


def test_leading_zero_and_jalali_date(client, env):
    from app.models.people import Applicant

    db = env["db"]
    a = db.query(Applicant).filter(Applicant.external_id == "T-XL-1").one()
    assert a.national_id == "0098765432"
    assert (a.birth_date.year, a.birth_date.month, a.birth_date.day) == (1991, 6, 11)


def test_excel_system_admin_role_rejected(client, env):
    from openpyxl import Workbook

    t = login(client, "t_admin")
    wb = Workbook()
    ws = wb.active
    ws.append(["username", "display_name", "email", "mobile", "role_codes",
               "unit_code", "counselor_code", "applicant_external_id"])
    ws.append(["t_xl_admin", "بد", None, None, "SYSTEM_ADMIN;RECEPTION", None, None, None])
    bio = io.BytesIO()
    wb.save(bio)
    r = client.post(f"{API}/imports/users/preview", headers=bearer(t),
                    files={"file": ("u.xlsx", bio.getvalue(), XLSX_MIME)})
    assert r.status_code == 200
    errs = " ".join(e["error"] for e in r.json()["errors"])
    assert "SYSTEM_ADMIN" in errs


# ---------- هم‌زمانی (بند ۱۱.۲) ----------

def test_case_version_conflict_409(client, env):
    t = login(client, "t_manager")
    cid = str(env["case"].id)
    ver = client.get(f"{API}/cases/{cid}", headers=bearer(t)).json()["version"]
    r2 = client.patch(f"{API}/cases/{cid}", headers=bearer(t),
                      json={"expected_version": ver + 5, "service_type": "X"})
    assert r2.status_code == 409
    r3 = client.patch(f"{API}/cases/{cid}", headers=bearer(t),
                      json={"expected_version": ver, "service_type": "مشاوره فردی"})
    assert r3.status_code == 200
    assert r3.json()["version"] == ver + 1


# ---------- رخداد و چاپ (بند ۱۴) ----------

def test_assessment_print_audit(client, env):
    from sqlalchemy import select

    from app.models.audit import AuditEvent

    ta = login(client, "t_counsa")
    cid = str(env["case"].id)
    r = client.post(f"{API}/cases/{cid}/assessments", headers=bearer(ta), json={"answers": {}})
    assert r.status_code == 201
    aid = r.json()["id"]
    ts = login(client, "t_super")
    rp = client.post(f"{API}/assessments/{aid}/print", headers=bearer(ts))
    assert rp.status_code == 204
    db = env["db"]
    ev = db.scalar(
        select(AuditEvent).where(AuditEvent.action == "assessment.printed")
        .order_by(AuditEvent.created_at.desc())
    )
    assert ev is not None and str(ev.entity_id) == aid


def test_audit_events_visibility(client, env):
    ta = login(client, "t_admin")
    assert client.get(f"{API}/audit-events", headers=bearer(ta)).status_code == 200
    tm = login(client, "t_manager")
    r2 = client.get(f"{API}/audit-events", headers=bearer(tm))
    assert r2.status_code == 200
    assert all(i["detail"] is None for i in r2.json()["items"])
    tr = login(client, "t_reception")
    assert client.get(f"{API}/audit-events", headers=bearer(tr)).status_code == 403


def test_reports_role_views(client, env):
    tm = login(client, "t_manager")
    r = client.get(f"{API}/reports/summary?days=30", headers=bearer(tm))
    assert r.status_code == 200 and r.json()["view"] == "staff"
    ta = login(client, "t_counsa")
    r2 = client.get(f"{API}/reports/summary?days=30", headers=bearer(ta))
    assert r2.status_code == 200 and r2.json()["view"] == "self"
