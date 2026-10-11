import io
import re
import uuid
from datetime import date
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.dates import normalize_digits, parse_date_cell
from app.models.people import Applicant, Counselor, OrganizationUnit, ServiceLocation, UserScope
from app.models.user import ROLE_CODES, Role, User, UserStatus

MAX_ROWS = 5000
MAX_FILE_SIZE = 5 * 1024 * 1024

USER_COLUMNS = ["username", "display_name", "email", "mobile", "role_codes", "unit_code", "counselor_code", "applicant_external_id"]
COUNSELOR_COLUMNS = ["counselor_code", "first_name", "last_name", "national_id", "professional_title", "degree", "specialty", "license_number", "license_expiry_date", "work_mobile", "email", "service_location_code", "capacity", "create_account", "username"]
APPLICANT_COLUMNS = ["external_id", "employee_code", "first_name", "last_name", "national_id", "birth_date", "gender", "marital_status", "education", "occupation", "mobile", "phone", "unit_code", "service_location_code", "province_code", "county_code", "city", "address", "referral_source", "referral_description", "create_account", "username"]

COLUMNS = {"users": USER_COLUMNS, "counselors": COUNSELOR_COLUMNS, "applicants": APPLICANT_COLUMNS}
IMPORT_TYPE_ROLES = {
    "users": {"SYSTEM_ADMIN"},
    "counselors": {"SYSTEM_ADMIN", "ORG_MANAGER"},
    "applicants": {"SYSTEM_ADMIN", "ORG_MANAGER", "RECEPTION"},
}

# ستون‌های شناسه/تماس که باید قالب متن باشند (حفظ صفر ابتدایی - بند ۸ سند)
ID_TEXT_COLUMNS = {
    "users": {"mobile"},
    "counselors": {"counselor_code", "national_id", "work_mobile"},
    "applicants": {"external_id", "employee_code", "national_id", "mobile", "phone"},
}

USERNAME_RE = re.compile(r"^[a-z0-9._-]{3,100}$")
TRUTHY = {"بله", "true", "1", "yes", "y"}

DESCRIPTIONS = {
    "users": [
        ("username", "الزامی و یکتا", "حروف انگلیسی/رقم/._- ؛ حداقل ۳ نویسه"),
        ("display_name", "الزامی", "نام نمایشی"),
        ("email", "اختیاری", "ایمیل"),
        ("mobile", "اختیاری", "قالب متن؛ نمونه: 09121234567"),
        ("role_codes", "الزامی", "جداشده با ; از: ORG_MANAGER;RECEPTION;COUNSELOR;TECHNICAL_SUPERVISOR;APPLICANT (SYSTEM_ADMIN مجاز نیست)"),
        ("unit_code", "اختیاری", "کد واحد سازمانی موجود"),
        ("counselor_code", "اختیاری", "اتصال حساب به مشاور موجود"),
        ("applicant_external_id", "اختیاری", "اتصال حساب به متقاضی موجود"),
    ],
    "counselors": [
        ("counselor_code", "الزامی و یکتا", "کد یکتای مشاور - قالب متن"),
        ("first_name", "الزامی", "نام"),
        ("last_name", "الزامی", "نام خانوادگی"),
        ("national_id", "اختیاری", "کد ملی ۱۰ رقمی - قالب متن"),
        ("professional_title", "اختیاری", "عنوان حرفه‌ای"),
        ("degree", "اختیاری", "مدرک"),
        ("specialty", "اختیاری", "تخصص"),
        ("license_number", "اختیاری", "شماره مجوز - قالب متن"),
        ("license_expiry_date", "اختیاری", "تاریخ اعتبار مجوز: 1403/05/12 شمسی یا 2025-08-12 میلادی"),
        ("work_mobile", "اختیاری", "تماس کاری - قالب متن"),
        ("email", "اختیاری", "ایمیل"),
        ("service_location_code", "اختیاری", "کد محل خدمت موجود"),
        ("capacity", "اختیاری", "ظرفیت پذیرش - عدد مثبت"),
        ("create_account", "اختیاری", "بله یا خیر - ایجاد حساب فقط نقش COUNSELOR می‌دهد"),
        ("username", "شرطی", "در صورت create_account=بله الزامی"),
    ],
    "applicants": [
        ("external_id", "اختیاری", "شناسه مبدأ یا جایگزین - قالب متن"),
        ("employee_code", "اختیاری", "شماره پرسنلی - قالب متن"),
        ("first_name", "الزامی", "نام"),
        ("last_name", "الزامی", "نام خانوادگی"),
        ("national_id", "اختیاری", "کد ملی ۱۰ رقمی - قالب متن"),
        ("birth_date", "اختیاری", "تاریخ تولد: 1370/03/21 شمسی یا 1991-06-11 میلادی"),
        ("gender", "اختیاری", "مرد یا زن"),
        ("marital_status", "اختیاری", "وضعیت تأهل"),
        ("education", "اختیاری", "تحصیلات"),
        ("occupation", "اختیاری", "شغل"),
        ("mobile", "اختیاری", "موبایل - قالب متن"),
        ("phone", "اختیاری", "تلفن - قالب متن"),
        ("unit_code", "اختیاری", "کد واحد سازمانی موجود"),
        ("service_location_code", "اختیاری", "کد محل خدمت موجود"),
        ("province_code", "اختیاری", "کد استان"),
        ("county_code", "اختیاری", "کد شهرستان"),
        ("city", "اختیاری", "شهر"),
        ("address", "اختیاری", "نشانی"),
        ("referral_source", "اختیاری", "منبع ارجاع"),
        ("referral_description", "اختیاری", "توضیح ارجاع"),
        ("create_account", "اختیاری", "بله یا خیر - حساب ساخته‌شده فقط نقش APPLICANT دارد"),
        ("username", "شرطی", "در صورت create_account=بله الزامی"),
    ],
}

RULES = [
    "فایل باید xlsx و بدون ماکرو باشد؛ حداکثر ۵ مگابایت و ۵۰۰۰ ردیف.",
    "ردیف اول شیت «داده» (عنوان ستون‌ها) حذف یا جابه‌جا نشود.",
    "ستون‌های شناسه و تماس حتماً قالب «متن» باشند تا صفر ابتدایی حفظ شود.",
    "تاریخ: شمسی با قالب YYYY/MM/DD (نمونه 1403/05/12) یا میلادی YYYY-MM-DD.",
    "در سلول‌ها فرمول ننویسید؛ سلول فرمول پذیرفته نمی‌شود.",
    "create_account فقط مقادیر «بله» یا «خیر».",
    "هیچ رمز عبوری در فایل وجود ندارد؛ حساب‌های جدید «در انتظار فعال‌سازی» هستند.",
    "در هر ردیف متقاضی حداقل یکی از شناسه‌ها (external_id / employee_code / national_id) الزامی است.",
]


def build_template(import_type: str) -> bytes:
    columns = COLUMNS[import_type]
    wb = Workbook()
    ws = wb.active
    ws.title = "داده"
    ws.append(columns)
    bold = Font(bold=True)
    for cell in ws[1]:
        cell.font = bold
    for idx in range(1, len(columns) + 1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = 22
    guide = wb.create_sheet("راهنما")
    guide.column_dimensions["A"].width = 28
    guide.column_dimensions["B"].width = 16
    guide.column_dimensions["C"].width = 80
    guide.append([f"قالب ورود {import_type} — سامانه Rosepsy (RS-PS-04-01)"])
    guide.append([])
    guide.append(["قواعد کلی:"])
    for rule in RULES:
        guide.append(["", "", rule])
    guide.append([])
    guide.append(["ستون", "الزام", "توضیح"])
    for name, req, desc in DESCRIPTIONS[import_type]:
        guide.append([name, req, desc])
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def parse_file(data: bytes, import_type: str) -> tuple[list[dict], str | None]:
    """بررسی سطح فایل و استخراج ردیف‌ها (بند ۸.۵ سند)."""
    columns = COLUMNS[import_type]
    try:
        wb = load_workbook(io.BytesIO(data), data_only=False, read_only=False)
    except Exception:
        return [], "فایل قابل خواندن نیست؛ فقط xlsx معتبر است"
    if not (1 <= len(wb.sheetnames) <= 5):
        return [], "تعداد شیت‌های فایل مجاز نیست (حداکثر ۵ شیت)"
    ws = wb.worksheets[0]
    if ws.max_row and ws.max_row > MAX_ROWS + 1:
        return [], f"تعداد ردیف‌ها بیش از سقف {MAX_ROWS} است"
    headers: list[str] = []
    for cell in ws[1]:
        v = cell.value
        headers.append(normalize_digits(str(v)).strip() if v is not None else "")
    if headers[: len(columns)] != columns:
        return [], "عنوان ستون‌ها با قالب مصوب مطابقت ندارد؛ قالب نمونه را دانلود کنید"

    id_cols = ID_TEXT_COLUMNS.get(import_type, set())
    rows: list[dict] = []
    for r in range(2, (ws.max_row or 1) + 1):
        values: dict[str, Any] = {}
        has_formula = False
        type_error: str | None = None
        empty = True
        for c in range(1, len(columns) + 1):
            cell = ws.cell(row=r, column=c)
            v = cell.value
            name = columns[c - 1]
            if v is None or (isinstance(v, str) and not v.strip()):
                values[name] = None
                continue
            empty = False
            if cell.data_type == "f" or (isinstance(v, str) and v.startswith("=")):
                has_formula = True
                values[name] = None
                continue
            if name in id_cols and isinstance(v, (int, float)):
                type_error = f"ستون «{name}» در ردیف {r} باید قالب متن باشد (صفر ابتدایی حفظ شود)"
                values[name] = None
                continue
            if isinstance(v, str):
                v = normalize_digits(v).strip() or None
            values[name] = v
        if empty:
            continue
        rows.append({"row_number": r, "has_formula": has_formula, "type_error": type_error, "values": values})
    return rows, None


def _as_str(v: Any) -> str | None:
    if v is None:
        return None
    s = normalize_digits(str(v)).strip()
    return s or None


def _parse_bool(v: Any) -> bool:
    if v is None:
        return False
    return normalize_digits(str(v)).strip().lower() in TRUTHY


def _check_mobile(errors: list[str], value: str | None, label: str) -> None:
    if value and not (value.isdigit() and 8 <= len(value) <= 15):
        errors.append(f"{label} باید فقط رقم و بین ۸ تا ۱۵ نویسه باشد")


def validate_row(db: Session, import_type: str, values: dict, seen: dict) -> dict:
    """اعتبارسنجی یک ردیف در برابر DB - در preview و commit دوباره اجرا می‌شود (بند ۸.۵)."""
    if import_type == "users":
        return _validate_users(db, values, seen)
    if import_type == "counselors":
        return _validate_counselors(db, values, seen)
    return _validate_applicants(db, values, seen)


def _validate_users(db: Session, values: dict, seen: dict) -> dict:
    errors: list[str] = []
    username = _as_str(values.get("username"))
    if username:
        username = username.lower().replace(" ", "_")
        if not USERNAME_RE.fullmatch(username):
            errors.append("نام کاربری فقط حروف انگلیسی/رقم/._- و حداقل ۳ نویسه")
    else:
        errors.append("username الزامی است")
    display_name = _as_str(values.get("display_name"))
    if not display_name:
        errors.append("display_name الزامی است")
    raw_roles = _as_str(values.get("role_codes")) or ""
    roles = sorted({r.strip().upper() for r in raw_roles.replace("؛", ";").split(";") if r.strip()})
    if not roles:
        errors.append("حداقل یک نقش لازم است")
    unknown = [r for r in roles if r not in ROLE_CODES]
    if unknown:
        errors.append(f"نقش نامعتبر: {','.join(unknown)}")
    if "SYSTEM_ADMIN" in roles:
        errors.append("ایجاد نقش SYSTEM_ADMIN از اکسل مجاز نیست")
    mobile = _as_str(values.get("mobile"))
    _check_mobile(errors, mobile, "موبایل")
    email = _as_str(values.get("email"))

    unit_id = counselor_id = applicant_id = None
    unit_code = _as_str(values.get("unit_code"))
    if unit_code:
        unit = db.scalar(select(OrganizationUnit).where(OrganizationUnit.code == unit_code))
        if unit is None:
            errors.append(f"واحد سازمانی «{unit_code}» یافت نشد")
        else:
            unit_id = str(unit.id)
    c_code = _as_str(values.get("counselor_code"))
    if c_code:
        c = db.scalar(select(Counselor).where(Counselor.counselor_code == c_code))
        if c is None:
            errors.append(f"کد مشاور «{c_code}» یافت نشد")
        else:
            counselor_id = str(c.id)
    a_ext = _as_str(values.get("applicant_external_id"))
    if a_ext:
        a = db.scalar(select(Applicant).where(Applicant.external_id == a_ext))
        if a is None:
            errors.append(f"شناسه متقاضی «{a_ext}» یافت نشد")
        else:
            applicant_id = str(a.id)

    existing = None
    will_update: list[str] = []
    if username and not errors:
        usernames = seen.setdefault("usernames", set())
        if username in usernames:
            errors.append("نام کاربری تکراری در فایل")
        else:
            usernames.add(username)
            u = db.scalar(select(User).where(User.username == username))
            if u is not None:
                existing = str(u.id)
                for f, nv in (("display_name", display_name), ("email", email), ("mobile", mobile)):
                    if nv is not None and nv != (getattr(u, f) or None):
                        will_update.append(f)

    clean = {"username": username, "display_name": display_name, "email": email, "mobile": mobile,
             "roles": roles, "unit_id": unit_id, "counselor_id": counselor_id, "applicant_id": applicant_id}
    return {"errors": errors, "existing_id": existing, "will_update": will_update, "clean": clean}


def _validate_counselors(db: Session, values: dict, seen: dict) -> dict:
    errors: list[str] = []
    code = _as_str(values.get("counselor_code"))
    if not code:
        errors.append("counselor_code الزامی است")
    first = _as_str(values.get("first_name"))
    last = _as_str(values.get("last_name"))
    if not first:
        errors.append("first_name الزامی است")
    if not last:
        errors.append("last_name الزامی است")
    national_id = _as_str(values.get("national_id"))
    if national_id and not (national_id.isdigit() and len(national_id) == 10):
        errors.append("کد ملی باید ۱۰ رقم متن باشد")
    expiry = None
    raw_expiry = values.get("license_expiry_date")
    if raw_expiry not in (None, ""):
        try:
            expiry = parse_date_cell(raw_expiry).isoformat()
        except ValueError as e:
            errors.append(str(e))
    work_mobile = _as_str(values.get("work_mobile"))
    _check_mobile(errors, work_mobile, "تماس کاری")
    capacity = None
    raw_cap = _as_str(values.get("capacity"))
    if raw_cap:
        try:
            capacity = int(raw_cap)
            if capacity <= 0:
                raise ValueError
        except ValueError:
            errors.append("ظرفیت باید عدد مثبت باشد")
            capacity = None
    loc_id = None
    loc_code = _as_str(values.get("service_location_code"))
    if loc_code:
        loc = db.scalar(select(ServiceLocation).where(ServiceLocation.code == loc_code))
        if loc is None:
            errors.append(f"محل خدمت «{loc_code}» یافت نشد")
        else:
            loc_id = str(loc.id)
    create_account = _parse_bool(values.get("create_account"))
    username = _as_str(values.get("username"))
    if create_account:
        if not username:
            errors.append("برای ایجاد حساب، username الزامی است")
        elif not USERNAME_RE.fullmatch(username.lower()):
            errors.append("نام کاربری حساب نامعتبر است")
        else:
            username = username.lower()

    existing = None
    will_update: list[str] = []
    if not errors:
        codes = seen.setdefault("counselor_codes", set())
        nids = seen.setdefault("counselor_nids", set())
        if code and code in codes:
            errors.append("کد مشاور تکراری در فایل")
        elif code:
            codes.add(code)
        if national_id and national_id in nids:
            errors.append("کد ملی تکراری در فایل")
        elif national_id:
            nids.add(national_id)
        if create_account and username and not errors:
            unames = seen.setdefault("usernames", set())
            if username in unames:
                errors.append("نام کاربری تکراری در فایل")
            else:
                unames.add(username)
                if db.scalar(select(User).where(User.username == username)):
                    errors.append("نام کاربری قبلاً در سامانه ثبت شده است")
        if not errors and code:
            c = db.scalar(select(Counselor).where(Counselor.counselor_code == code))
            if c is not None:
                existing = str(c.id)
        if not errors and existing is None and national_id:
            c = db.scalar(select(Counselor).where(Counselor.national_id == national_id))
            if c is not None:
                existing = str(c.id)
        if existing:
            c = db.get(Counselor, uuid.UUID(existing))
            for f, nv in (("professional_title", _as_str(values.get("professional_title"))),
                          ("degree", _as_str(values.get("degree"))),
                          ("specialty", _as_str(values.get("specialty"))),
                          ("license_number", _as_str(values.get("license_number"))),
                          ("work_mobile", work_mobile), ("email", _as_str(values.get("email")))):
                if nv is not None and nv != (getattr(c, f) or None):
                    will_update.append(f)
            if capacity is not None and capacity != (c.capacity or 0):
                will_update.append("capacity")

    clean = {"counselor_code": code, "first_name": first, "last_name": last, "national_id": national_id,
             "professional_title": _as_str(values.get("professional_title")), "degree": _as_str(values.get("degree")),
             "specialty": _as_str(values.get("specialty")), "license_number": _as_str(values.get("license_number")),
             "license_expiry_date": expiry, "work_mobile": work_mobile, "email": _as_str(values.get("email")),
             "service_location_id": loc_id, "capacity": capacity, "create_account": create_account, "username": username}
    return {"errors": errors, "existing_id": existing, "will_update": will_update, "clean": clean}


def _validate_applicants(db: Session, values: dict, seen: dict) -> dict:
    errors: list[str] = []
    first = _as_str(values.get("first_name"))
    last = _as_str(values.get("last_name"))
    if not first:
        errors.append("first_name الزامی است")
    if not last:
        errors.append("last_name الزامی است")
    external_id = _as_str(values.get("external_id"))
    employee_code = _as_str(values.get("employee_code"))
    national_id = _as_str(values.get("national_id"))
    if not any([external_id, employee_code, national_id]):
        errors.append("حداقل یکی از شناسه‌ها (external_id / employee_code / national_id) لازم است")
    if national_id and not (national_id.isdigit() and len(national_id) == 10):
        errors.append("کد ملی باید ۱۰ رقم متن باشد")
    if employee_code and not employee_code.isdigit():
        errors.append("شماره پرسنلی باید فقط رقم و متن باشد")
    birth = None
    raw_birth = values.get("birth_date")
    if raw_birth not in (None, ""):
        try:
            birth = parse_date_cell(raw_birth).isoformat()
        except ValueError as e:
            errors.append(str(e))
    mobile = _as_str(values.get("mobile"))
    phone = _as_str(values.get("phone"))
    _check_mobile(errors, mobile, "موبایل")
    _check_mobile(errors, phone, "تلفن")
    g = _as_str(values.get("gender"))
    gmap = {"مرد": "male", "زن": "female", "male": "male", "female": "female", "m": "male", "f": "female"}
    gender = (gmap.get(g.lower(), g.lower()) if g else None)
    unit_id = loc_id = None
    unit_code = _as_str(values.get("unit_code"))
    if unit_code:
        unit = db.scalar(select(OrganizationUnit).where(OrganizationUnit.code == unit_code))
        if unit is None:
            errors.append(f"واحد سازمانی «{unit_code}» یافت نشد")
        else:
            unit_id = str(unit.id)
    loc_code = _as_str(values.get("service_location_code"))
    if loc_code:
        loc = db.scalar(select(ServiceLocation).where(ServiceLocation.code == loc_code))
        if loc is None:
            errors.append(f"محل خدمت «{loc_code}» یافت نشد")
        else:
            loc_id = str(loc.id)
    create_account = _parse_bool(values.get("create_account"))
    username = _as_str(values.get("username"))
    if create_account:
        if not username:
            errors.append("برای ایجاد حساب، username الزامی است")
        elif not USERNAME_RE.fullmatch(username.lower()):
            errors.append("نام کاربری حساب نامعتبر است")
        else:
            username = username.lower()

    existing = None
    will_update: list[str] = []
    if not errors:
        s_ext = seen.setdefault("applicant_ext", set())
        s_emp = seen.setdefault("applicant_emp", set())
        s_nid = seen.setdefault("applicant_nid", set())
        if external_id and external_id in s_ext:
            errors.append("external_id تکراری در فایل")
        elif external_id:
            s_ext.add(external_id)
        if employee_code and employee_code in s_emp:
            errors.append("شماره پرسنلی تکراری در فایل")
        elif employee_code:
            s_emp.add(employee_code)
        if national_id and national_id in s_nid:
            errors.append("کد ملی تکراری در فایل")
        elif national_id:
            s_nid.add(national_id)
        if create_account and username and not errors:
            unames = seen.setdefault("usernames", set())
            if username in unames:
                errors.append("نام کاربری تکراری در فایل")
            else:
                unames.add(username)
                if db.scalar(select(User).where(User.username == username)):
                    errors.append("نام کاربری قبلاً در سامانه ثبت شده است")
        if not errors:
            for val, col in ((external_id, Applicant.external_id), (employee_code, Applicant.employee_code), (national_id, Applicant.national_id)):
                if val:
                    a = db.scalar(select(Applicant).where(col == val))
                    if a is not None:
                        existing = str(a.id)
                        break
        if existing:
            a = db.get(Applicant, uuid.UUID(existing))
            for f, nv in (("mobile", mobile), ("phone", phone),
                          ("education", _as_str(values.get("education"))),
                          ("occupation", _as_str(values.get("occupation"))),
                          ("marital_status", _as_str(values.get("marital_status"))),
                          ("city", _as_str(values.get("city"))),
                          ("address", _as_str(values.get("address"))),
                          ("referral_source", _as_str(values.get("referral_source"))),
                          ("referral_description", _as_str(values.get("referral_description")))):
                if nv is not None and nv != (getattr(a, f) or None):
                    will_update.append(f)
            if unit_id and unit_id != (str(a.unit_id) if a.unit_id else None):
                will_update.append("unit")
            if loc_id and loc_id != (str(a.service_location_id) if a.service_location_id else None):
                will_update.append("service_location")

    clean = {"external_id": external_id, "employee_code": employee_code, "national_id": national_id,
             "first_name": first, "last_name": last, "birth_date": birth, "gender": gender,
             "marital_status": _as_str(values.get("marital_status")), "education": _as_str(values.get("education")),
             "occupation": _as_str(values.get("occupation")), "mobile": mobile, "phone": phone,
             "unit_id": unit_id, "service_location_id": loc_id,
             "province_code": _as_str(values.get("province_code")), "county_code": _as_str(values.get("county_code")),
             "city": _as_str(values.get("city")), "address": _as_str(values.get("address")),
             "referral_source": _as_str(values.get("referral_source")),
             "referral_description": _as_str(values.get("referral_description")),
             "create_account": create_account, "username": username}
    return {"errors": errors, "existing_id": existing, "will_update": will_update, "clean": clean}


# ---------- ایجاد رکوردها (پروفایل و حساب در یک تراکنش - بند ۸.۵) ----------

def create_records(db: Session, import_type: str, clean: dict, actor_roles: set[str]) -> None:
    if import_type == "users":
        _create_user_row(db, clean)
    elif import_type == "counselors":
        # بند ۸.۳: ورودکننده فاقد مجوز ساخت حساب، فقط پروفایل مشاور ایجاد می‌کند
        allow_account = bool(clean.get("create_account")) and "SYSTEM_ADMIN" in actor_roles
        _create_counselor_row(db, clean, allow_account)
    else:
        _create_applicant_row(db, clean, bool(clean.get("create_account")))


def _create_user_row(db: Session, clean: dict) -> None:
    user = User(
        username=clean["username"],
        display_name=clean.get("display_name") or clean["username"],
        email=clean.get("email"),
        mobile=clean.get("mobile"),
        status=UserStatus.PENDING,
        must_change_password=True,
    )
    roles = db.scalars(select(Role).where(Role.code.in_(clean["roles"]))).all()
    user.roles = list(roles)
    db.add(user)
    db.flush()
    if clean.get("unit_id"):
        db.add(UserScope(user_id=user.id, unit_id=uuid.UUID(clean["unit_id"])))
    if clean.get("counselor_id"):
        c = db.get(Counselor, uuid.UUID(clean["counselor_id"]))
        if c.user_id is not None:
            raise ValueError("مشاور موردنظر قبلاً حساب مرتبط دارد")
        c.user_id = user.id
    if clean.get("applicant_id"):
        a = db.get(Applicant, uuid.UUID(clean["applicant_id"]))
        if a.user_id is not None:
            raise ValueError("متقاضی موردنظر قبلاً حساب مرتبط دارد")
        a.user_id = user.id


def _create_counselor_row(db: Session, clean: dict, create_account: bool) -> None:
    c = Counselor(
        counselor_code=clean["counselor_code"],
        first_name=clean["first_name"],
        last_name=clean["last_name"],
        national_id=clean.get("national_id"),
        professional_title=clean.get("professional_title"),
        degree=clean.get("degree"),
        specialty=clean.get("specialty"),
        license_number=clean.get("license_number"),
        license_expiry_date=date.fromisoformat(clean["license_expiry_date"]) if clean.get("license_expiry_date") else None,
        work_mobile=clean.get("work_mobile"),
        email=clean.get("email"),
        service_location_id=uuid.UUID(clean["service_location_id"]) if clean.get("service_location_id") else None,
        capacity=clean.get("capacity"),
        status="active",
    )
    db.add(c)
    db.flush()
    if create_account:
        user = User(
            username=clean["username"],
            display_name=f"{clean['first_name']} {clean['last_name']}",
            status=UserStatus.PENDING,
            must_change_password=True,
        )
        user.roles.append(db.get(Role, "COUNSELOR"))
        db.add(user)
        db.flush()
        c.user_id = user.id


def _create_applicant_row(db: Session, clean: dict, create_account: bool) -> None:
    a = Applicant(
        external_id=clean.get("external_id"),
        employee_code=clean.get("employee_code"),
        national_id=clean.get("national_id"),
        first_name=clean["first_name"],
        last_name=clean["last_name"],
        birth_date=date.fromisoformat(clean["birth_date"]) if clean.get("birth_date") else None,
        gender=clean.get("gender"),
        marital_status=clean.get("marital_status"),
        education=clean.get("education"),
        occupation=clean.get("occupation"),
        mobile=clean.get("mobile"),
        phone=clean.get("phone"),
        unit_id=uuid.UUID(clean["unit_id"]) if clean.get("unit_id") else None,
        service_location_id=uuid.UUID(clean["service_location_id"]) if clean.get("service_location_id") else None,
        province_code=clean.get("province_code"),
        county_code=clean.get("county_code"),
        city=clean.get("city"),
        address=clean.get("address"),
        referral_source=clean.get("referral_source"),
        referral_description=clean.get("referral_description"),
        status="active",
        record_source="excel",
    )
    db.add(a)
    db.flush()
    if create_account:
        user = User(
            username=clean["username"],
            display_name=f"{clean['first_name']} {clean['last_name']}",
            status=UserStatus.PENDING,
            must_change_password=True,
        )
        user.roles.append(db.get(Role, "APPLICANT"))
        db.add(user)
        db.flush()
        a.user_id = user.id


def apply_updates(db: Session, import_type: str, existing_id: str, clean: dict) -> None:
    """به‌روزرسانی فقط روی فیلدهای مجاز (بند ۸.۵: تغییر نقش/وضعیت/پیوند هویتی از اکسل ممنوع)."""
    if import_type == "users":
        u = db.get(User, uuid.UUID(existing_id))
        for f in ("display_name", "email", "mobile"):
            if clean.get(f) is not None:
                setattr(u, f, clean[f])
    elif import_type == "counselors":
        c = db.get(Counselor, uuid.UUID(existing_id))
        for f in ("first_name", "last_name", "professional_title", "degree", "specialty", "license_number", "work_mobile", "email"):
            if clean.get(f) is not None:
                setattr(c, f, clean[f])
        if clean.get("license_expiry_date"):
            c.license_expiry_date = date.fromisoformat(clean["license_expiry_date"])
        if clean.get("capacity") is not None:
            c.capacity = clean["capacity"]
        if clean.get("service_location_id"):
            c.service_location_id = uuid.UUID(clean["service_location_id"])
    else:
        a = db.get(Applicant, uuid.UUID(existing_id))
        for f in ("first_name", "last_name", "gender", "marital_status", "education", "occupation",
                  "mobile", "phone", "city", "address", "referral_source", "referral_description"):
            if clean.get(f) is not None:
                setattr(a, f, clean[f])
        if clean.get("birth_date"):
            a.birth_date = date.fromisoformat(clean["birth_date"])
        if clean.get("unit_id"):
            a.unit_id = uuid.UUID(clean["unit_id"])
        if clean.get("service_location_id"):
            a.service_location_id = uuid.UUID(clean["service_location_id"])
